"""Agent tests with a fake OpenRouter: no real AI is called."""

from __future__ import annotations

import io

import httpx
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.models import AICall, EvidenceFile
from tests.conftest import ADMIN_EMAIL, FakeOpenRouter, login, project_id


def _new_project(client: TestClient) -> str:
    login(client, "owner@dev.local")
    return client.post(
        "/api/projects",
        json={
            "name": "Invoice Bot",
            "business_unit": "Finance",
            "sponsor": "Ko Aung",
            "mode": "assess",
            "stage": "Build",
        },
    ).json()["id"]


def _calls(db: Session) -> list[AICall]:
    db.expire_all()
    return list(db.scalars(select(AICall).order_by(AICall.id)).all())


# ---- Interview ----


def test_interview_records_ai_answer(client: TestClient, fake_llm: FakeOpenRouter, db: Session):
    pid = _new_project(client)
    fake_llm.reply(
        {"answer": "Partly", "evidence": "HR answers link to the policy page.", "followUp": ""}
    )
    r = client.post(
        f"/api/projects/{pid}/agent/interview",
        json={"question_id": "5.4", "reply": "HR answers show a link. Payroll no."},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["answer"] == "Partly" and body["followUp"] == ""
    saved = body["saved"]
    assert saved["answer"]["source"] == "ai_interview"
    assert saved["answer"]["confirmed"] is False  # marked until a person confirms
    assert saved["score"]["verdict"] == "go"  # the rules decide, not the AI

    sent = fake_llm.requests[-1]
    assert sent["url"] == "https://openrouter.test/api/v1/chat/completions"
    assert sent["headers"]["authorization"] == "Bearer test-key"
    assert sent["json"]["model"] == "fast-model"
    assert sent["json"]["response_format"] == {"type": "json_object"}
    prompt = fake_llm.last_prompt
    assert (
        "Project: Invoice Bot. Business unit: Finance. Sponsor: Ko Aung. Owner: Test Owner."
        in prompt
    )
    assert "Purpose of this check: assessing an existing project." in prompt
    assert "Dimension 5: Quality & Assurance." in prompt
    assert "Question 5.4: Can users verify outputs and see the supporting sources?" in prompt
    assert 'The owner now replied: """HR answers show a link. Payroll no."""' in prompt
    assert "Earlier answer" not in prompt

    call = _calls(db)[-1]
    assert (call.feature, call.model, call.ok) == ("interview", "fast-model", True)
    assert call.prompt_tokens == 120 and call.completion_tokens == 30
    assert call.cost_usd == 0.0001 and call.project_id == pid and call.latency_ms is not None

    hist = client.get(f"/api/projects/{pid}/history").json()
    assert hist[0]["source"] == "ai_interview" and hist[0]["new_answer"] == "Partly"


def test_interview_dont_know(client: TestClient, fake_llm: FakeOpenRouter):
    pid = _new_project(client)
    fake_llm.reply({"answer": "Don't know", "evidence": "The owner does not know.", "followUp": ""})
    r = client.post(
        f"/api/projects/{pid}/agent/interview",
        json={"question_id": "8.1", "reply": "No idea, sorry."},
    )
    assert r.json()["saved"]["answer"]["answer"] == "Dont_know"


def test_vague_reply_gets_follow_up(client: TestClient, fake_llm: FakeOpenRouter, db: Session):
    pid = _new_project(client)
    fake_llm.reply({"answer": "", "evidence": "", "followUp": "Who set the quality target?"})
    r = client.post(
        f"/api/projects/{pid}/agent/interview", json={"question_id": "5.1", "reply": "Sort of."}
    )
    assert r.json()["followUp"] == "Who set the quality target?"
    assert r.json()["saved"] is None
    assert client.get(f"/api/projects/{pid}").json()["answers"] == []

    fake_llm.reply({"answer": "Yes", "evidence": "The CFO set 95%.", "followUp": ""})
    client.post(
        f"/api/projects/{pid}/agent/interview",
        json={"question_id": "5.1", "reply": "The CFO set 95%.", "context": "Sort of."},
    )
    assert 'Earlier answer from the owner on this question: """Sort of."""' in fake_llm.last_prompt


def test_invalid_json_saves_nothing(client: TestClient, fake_llm: FakeOpenRouter, db: Session):
    pid = _new_project(client)
    fake_llm.reply("Sure! The answer is yes.")
    r = client.post(
        f"/api/projects/{pid}/agent/interview", json={"question_id": "5.1", "reply": "Yes, 95%."}
    )
    assert r.status_code == 502
    assert r.json()["detail"].startswith("The AI reply could not be read.")
    assert client.get(f"/api/projects/{pid}").json()["answers"] == []
    call = _calls(db)[-1]
    assert not call.ok and call.error == "invalid_json"
    assert len(fake_llm.requests) == 1  # no automatic retry

    for bad in (
        {"answer": "Maybe", "evidence": "x", "followUp": ""},
        {"answer": "", "evidence": "", "followUp": ""},
    ):
        fake_llm.reply(bad)
        r = client.post(
            f"/api/projects/{pid}/agent/interview", json={"question_id": "5.1", "reply": "Yes."}
        )
        assert r.status_code == 502
        assert _calls(db)[-1].error == "invalid_schema"
    assert client.get(f"/api/projects/{pid}").json()["answers"] == []


def test_fenced_json_is_read(client: TestClient, fake_llm: FakeOpenRouter):
    pid = _new_project(client)
    fake_llm.reply('```json\n{"answer":"Yes","evidence":"Done.","followUp":""}\n```')
    r = client.post(
        f"/api/projects/{pid}/agent/interview", json={"question_id": "5.1", "reply": "Yes."}
    )
    assert r.json()["answer"] == "Yes"


def test_service_errors_are_logged(client: TestClient, fake_llm: FakeOpenRouter, db: Session):
    pid = _new_project(client)
    fake_llm.reply({}, status=500)
    r = client.post(
        f"/api/projects/{pid}/agent/interview", json={"question_id": "5.1", "reply": "Yes."}
    )
    assert r.status_code == 502 and "did not answer" in r.json()["detail"]
    fake_llm.fail(httpx.ReadTimeout("slow"))
    r = client.post(
        f"/api/projects/{pid}/agent/interview", json={"question_id": "5.1", "reply": "Yes."}
    )
    assert r.status_code == 504
    assert [c.error for c in _calls(db)] == ["http_500", "timeout"]


def test_not_set_up(client: TestClient, monkeypatch):
    from app.config import get_settings

    monkeypatch.setenv("OPENROUTER_API_KEY", "")
    get_settings.cache_clear()
    pid = _new_project(client)
    r = client.post(
        f"/api/projects/{pid}/agent/interview", json={"question_id": "5.1", "reply": "Yes."}
    )
    assert r.status_code == 503 and "not set up" in r.json()["detail"]


def test_interview_checks(client: TestClient, fake_llm: FakeOpenRouter):
    pid = _new_project(client)
    r = client.post(
        f"/api/projects/{pid}/agent/interview", json={"question_id": "9.9", "reply": "Yes."}
    )
    assert r.status_code == 404
    assert (
        client.post(
            f"/api/projects/{pid}/agent/interview", json={"question_id": "5.1", "reply": ""}
        ).status_code
        == 422
    )
    login(client, "approver@dev.local")
    assert (
        client.post(
            f"/api/projects/{pid}/agent/interview", json={"question_id": "5.1", "reply": "Yes."}
        ).status_code
        == 403
    )
    login(client, "owner@dev.local")
    other = project_id(client, "CFC Model")
    assert (
        client.post(
            f"/api/projects/{other}/agent/interview", json={"question_id": "5.1", "reply": "Yes."}
        ).status_code
        == 403
    )
    assert fake_llm.requests == []  # refused before any AI call


# ---- Undo ----


def test_undo_last_agent_answer(client: TestClient, fake_llm: FakeOpenRouter):
    pid = _new_project(client)
    client.put(f"/api/projects/{pid}/answers/5.4", json={"answer": "No", "evidence": "Mine."})
    fake_llm.reply({"answer": "Partly", "evidence": "Some links.", "followUp": ""})
    client.post(
        f"/api/projects/{pid}/agent/interview",
        json={"question_id": "5.4", "reply": "Some answers link."},
    )

    r = client.post(f"/api/projects/{pid}/answers/5.4/undo")
    assert r.status_code == 200
    a = r.json()["answer"]
    assert (a["answer"], a["evidence"], a["source"]) == ("No", "Mine.", "person")
    hist = client.get(f"/api/projects/{pid}/history").json()
    assert hist[0]["source"] == "undo"
    assert (hist[0]["old_answer"], hist[0]["new_answer"]) == ("Partly", "No")
    # Nothing left to undo.
    assert client.post(f"/api/projects/{pid}/answers/5.4/undo").status_code == 409


def test_undo_to_blank_and_not_after_a_person_edit(client: TestClient, fake_llm: FakeOpenRouter):
    pid = _new_project(client)
    fake_llm.reply({"answer": "Yes", "evidence": "A.", "followUp": ""})
    client.post(f"/api/projects/{pid}/agent/interview", json={"question_id": "1.1", "reply": "A"})
    r = client.post(f"/api/projects/{pid}/answers/1.1/undo")
    assert r.json()["answer"]["answer"] is None
    assert r.json()["score"]["verdict"] == "not_assessed"

    fake_llm.reply({"answer": "Yes", "evidence": "B.", "followUp": ""})
    client.post(f"/api/projects/{pid}/agent/interview", json={"question_id": "1.2", "reply": "B"})
    client.put(f"/api/projects/{pid}/answers/1.2", json={"answer": "Partly", "evidence": "B."})
    assert client.post(f"/api/projects/{pid}/answers/1.2/undo").status_code == 409


# ---- Read evidence ----

EVIDENCE_REPLY = [
    {"id": "3.2", "answer": "Yes", "evidence": "The team has read access to the ERP."},
    {"id": "5.1", "answer": "No", "evidence": "No quality target is set."},
    {"id": "9.9", "answer": "Yes", "evidence": "Unknown question."},
    {"id": "4.1", "answer": "Don't know", "evidence": "Not allowed here."},
    {"id": "3.2", "answer": "No", "evidence": "Duplicate."},
]


def test_read_pasted_text(client: TestClient, fake_llm: FakeOpenRouter, db: Session):
    pid = _new_project(client)
    client.put(f"/api/projects/{pid}/answers/2.1", json={"answer": "Yes"})
    fake_llm.reply(EVIDENCE_REPLY)
    r = client.post(
        f"/api/projects/{pid}/agent/evidence",
        data={"text": "We can read the ERP. No quality target yet."},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert [(s["question_id"], s["answer"]) for s in body["suggestions"]] == [
        ("3.2", "Yes"),
        ("5.1", "No"),
    ]
    assert body["filename"] == "Pasted text" and body["truncated"] is False
    assert all(s["status"] == "pending" for s in body["suggestions"])

    sent = fake_llm.requests[-1]["json"]
    assert sent["model"] == "default-model"
    prompt = fake_llm.last_prompt
    assert "2.1 Is there a named business owner? [already answered: Yes]" in prompt
    assert "2.2 Are responsibilities clear to the business unit or function?\n" in prompt
    assert '"""We can read the ERP. No quality target yet."""' in prompt
    assert prompt.count("\n1.") == 5 and "8.5 What routine" in prompt

    stored = db.scalars(select(EvidenceFile)).all()
    assert len(stored) == 1 and stored[0].text_extract.startswith("We can read")
    # Suggestions are not answers until a person accepts them.
    assert [a["question_id"] for a in client.get(f"/api/projects/{pid}").json()["answers"]] == [
        "2.1"
    ]
    assert _calls(db)[-1].feature == "evidence"


def test_accept_dismiss_and_accept_all(client: TestClient, fake_llm: FakeOpenRouter):
    pid = _new_project(client)
    fake_llm.reply(EVIDENCE_REPLY[:2] + [{"id": "7.1", "answer": "Partly", "evidence": "C."}])
    sugg = client.post(f"/api/projects/{pid}/agent/evidence", data={"text": "notes"}).json()[
        "suggestions"
    ]
    first, second, third = sugg

    r = client.post(f"/api/suggestions/{first['id']}/accept")
    saved = r.json()["saved"]
    assert saved["answer"]["source"] == "ai_evidence" and saved["answer"]["confirmed"] is True
    assert saved["score"]["verdict"] == "ready"
    assert client.post(f"/api/suggestions/{first['id']}/accept").status_code == 409

    assert (
        client.post(f"/api/suggestions/{second['id']}/dismiss").json()["suggestion"]["status"]
        == "dismissed"
    )
    pending = client.get(f"/api/projects/{pid}/suggestions").json()
    assert [s["id"] for s in pending] == [third["id"]]

    r = client.post(f"/api/projects/{pid}/suggestions/accept-all")
    assert r.json()["accepted"] == 1 and r.json()["score"]["answered"] == 2
    assert client.get(f"/api/projects/{pid}/suggestions").json() == []
    hist = client.get(f"/api/projects/{pid}/history").json()
    assert {h["source"] for h in hist} == {"ai_evidence"}

    login(client, "approver@dev.local")
    assert client.post(f"/api/projects/{pid}/suggestions/accept-all").status_code == 403


def test_new_reading_replaces_open_suggestions(client: TestClient, fake_llm: FakeOpenRouter):
    pid = _new_project(client)
    fake_llm.reply([{"id": "3.2", "answer": "Partly", "evidence": "Old."}])
    client.post(f"/api/projects/{pid}/agent/evidence", data={"text": "a"})
    fake_llm.reply([{"id": "3.2", "answer": "Yes", "evidence": "New."}])
    client.post(f"/api/projects/{pid}/agent/evidence", data={"text": "b"})
    pending = client.get(f"/api/projects/{pid}/suggestions").json()
    assert [(s["answer"], s["evidence"]) for s in pending] == [("Yes", "New.")]


def test_upload_docx_and_limits(client: TestClient, fake_llm: FakeOpenRouter, db: Session):
    from docx import Document

    pid = _new_project(client)
    doc = Document()
    doc.add_paragraph("Data owner: Finance team.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "System"
    table.cell(0, 1).text = "SAP"
    buf = io.BytesIO()
    doc.save(buf)
    fake_llm.reply([])
    r = client.post(
        f"/api/projects/{pid}/agent/evidence",
        files={"file": ("plan.docx", buf.getvalue(), "application/octet-stream")},
    )
    assert r.status_code == 200, r.text
    assert r.json()["filename"] == "plan.docx" and r.json()["suggestions"] == []
    assert "Data owner: Finance team." in fake_llm.last_prompt
    assert "System | SAP" in fake_llm.last_prompt
    stored = db.scalars(select(EvidenceFile)).one()
    assert stored.content == buf.getvalue()  # the file stays in our database

    # Long text: only the first 24,000 characters go to the AI.
    fake_llm.reply([])
    r = client.post(
        f"/api/projects/{pid}/agent/evidence",
        files={"file": ("long.txt", ("A" * 30_000).encode(), "text/plain")},
    )
    assert r.json()["truncated"] is True and r.json()["chars_read"] == 24_000
    assert "A" * 24_000 + '"""' in fake_llm.last_prompt

    r = client.post(
        f"/api/projects/{pid}/agent/evidence",
        files={"file": ("x.exe", b"MZ", "application/octet-stream")},
    )
    assert r.status_code == 422 and "cannot be read" in r.json()["detail"]
    r = client.post(
        f"/api/projects/{pid}/agent/evidence",
        files={"file": ("broken.pdf", b"not a pdf", "application/pdf")},
    )
    assert r.status_code == 422
    assert client.post(f"/api/projects/{pid}/agent/evidence", data={"text": " "}).status_code == 422


def test_ai_call_log_for_admin(client: TestClient, fake_llm: FakeOpenRouter):
    pid = _new_project(client)
    fake_llm.reply({"answer": "Yes", "evidence": "x", "followUp": ""})
    client.post(f"/api/projects/{pid}/agent/interview", json={"question_id": "1.1", "reply": "y"})
    assert client.get("/api/admin/ai-calls").status_code == 403
    login(client, ADMIN_EMAIL)
    rows = client.get("/api/admin/ai-calls").json()
    assert rows[0]["feature"] == "interview" and rows[0]["ok"] is True
    assert "prompt" not in rows[0]
