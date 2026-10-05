"""Decision brief, portfolio summary and export. A fake OpenRouter stands in for the AI."""

from __future__ import annotations

import csv
import io

import httpx
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.models import AICall, Brief
from app.services.export import COLUMNS
from tests.conftest import ADMIN_EMAIL, FakeOpenRouter, login, project_id

GOOD_BRIEF = {
    "headline": "Employee Assistant must fix its quality bar before the next stage.",
    "summary": "The owner and the steering committee are in place. The BU does not know yet.",
    "actions": [
        {"question": "5.1", "action": "Agree the quality target.", "owner": "Chief People Officer"},
        {"question": "2.2", "action": "Brief the business unit.", "owner": ""},
        {"question": "9.9", "action": "Unknown question, dropped.", "owner": "x"},
    ],
    "risks": ["Launch without a quality bar.", "", "Low adoption.", "Third.", "Fourth."],
}


def _calls(db: Session, feature: str) -> list[AICall]:
    db.expire_all()
    return list(db.scalars(select(AICall).where(AICall.feature == feature)).all())


# ---- Decision brief ----


def test_brief_keeps_the_rules_verdict(client: TestClient, fake_llm: FakeOpenRouter, db: Session):
    login(client, ADMIN_EMAIL)
    pid = project_id(client, "Employee Assistant")
    rules = client.get(f"/api/projects/{pid}").json()["score"]
    # The AI tries to set its own verdict and scores; they are ignored.
    fake_llm.reply({**GOOD_BRIEF, "verdict": "ready", "lowest": 5.0})
    r = client.post(f"/api/projects/{pid}/briefs")
    assert r.status_code == 201, r.text
    brief = r.json()
    assert brief["verdict"] == rules["verdict"] == "go"
    assert brief["lowest"] == rules["lowest"]
    assert brief["weakest"] == rules["weakest"]
    assert brief["answered"] == rules["answered"]
    assert [d["score"] for d in brief["dimensions"]] == [d["score"] for d in rules["dimensions"]]
    # Unknown question dropped; empty owner becomes "To assign"; at most 3 risks.
    assert [a["question_id"] for a in brief["actions"]] == ["5.1", "2.2"]
    assert brief["actions"][1]["owner"] == "To assign"
    assert brief["risks"] == ["Launch without a quality bar.", "Low adoption.", "Third."]
    assert brief["changed_since"] is False

    # The prompt holds the rules result and only this project's answers.
    prompt = fake_llm.requests[0]["json"]["messages"][0]["content"]
    assert 'verdict "Go with actions"' in prompt
    assert "Answered 8 of 40 questions." in prompt
    assert "5.1 " in prompt and "→ No (Business owner has not set the quality %.)" in prompt
    assert "Open questions: 1.1, 1.2" in prompt
    assert "CFC" not in prompt
    assert fake_llm.requests[0]["json"]["model"] == "default-model"
    assert len(_calls(db, "brief")) == 1


def test_brief_marks_later_changes_and_keeps_history(client: TestClient, fake_llm: FakeOpenRouter):
    login(client, ADMIN_EMAIL)
    pid = project_id(client, "Employee Assistant")
    fake_llm.reply(GOOD_BRIEF)
    first = client.post(f"/api/projects/{pid}/briefs").json()
    # Three more No answers: the rules move the verdict; the old brief keeps its snapshot.
    for qid in ("2.1", "2.3", "2.4"):
        client.put(f"/api/projects/{pid}/answers/{qid}", json={"answer": "No", "evidence": ""})
    briefs = client.get(f"/api/projects/{pid}/briefs").json()
    assert [b["id"] for b in briefs] == [first["id"]]
    assert briefs[0]["verdict"] == "go" and briefs[0]["changed_since"] is True

    fake_llm.reply(GOOD_BRIEF)
    second = client.post(f"/api/projects/{pid}/briefs").json()
    now = client.get(f"/api/projects/{pid}").json()["score"]
    assert second["verdict"] == now["verdict"] == "fix"
    assert second["changed_since"] is False
    briefs = client.get(f"/api/projects/{pid}/briefs").json()
    assert [b["id"] for b in briefs] == [second["id"], first["id"]]  # latest first


def test_brief_bad_reply_saves_nothing(client: TestClient, fake_llm: FakeOpenRouter, db: Session):
    login(client, ADMIN_EMAIL)
    pid = project_id(client, "Employee Assistant")
    fake_llm.reply("not json at all")
    r = client.post(f"/api/projects/{pid}/briefs")
    assert r.status_code == 502
    assert "could not be read" in r.json()["detail"]
    fake_llm.reply({"headline": "", "summary": "x", "actions": [], "risks": []})
    assert client.post(f"/api/projects/{pid}/briefs").status_code == 502
    db.expire_all()
    assert db.scalars(select(Brief)).all() == []
    assert [c.error for c in _calls(db, "brief")] == ["invalid_json", "invalid_schema"]


def test_brief_needs_answers_and_edit_rights(client: TestClient, fake_llm: FakeOpenRouter):
    login(client, ADMIN_EMAIL)
    empty = project_id(client, "City Care Agent")
    r = client.post(f"/api/projects/{empty}/briefs")
    assert r.status_code == 409
    assert "at least one question" in r.json()["detail"]

    pid = project_id(client, "Employee Assistant")
    login(client, "approver@dev.local")
    assert client.post(f"/api/projects/{pid}/briefs").status_code == 403
    login(client, "owner@dev.local")  # not the owner of this project
    assert client.post(f"/api/projects/{pid}/briefs").status_code == 403
    assert fake_llm.requests == []


def test_only_approvers_can_approve(client: TestClient, fake_llm: FakeOpenRouter):
    login(client, ADMIN_EMAIL)
    pid = project_id(client, "Employee Assistant")
    fake_llm.reply(GOOD_BRIEF)
    bid = client.post(f"/api/projects/{pid}/briefs").json()["id"]

    for email in (ADMIN_EMAIL, "reviewer@dev.local", "owner@dev.local"):
        login(client, email)
        r = client.post(f"/api/briefs/{bid}/approve")
        assert r.status_code == 403, email
        assert r.json()["detail"] == "Only approvers can approve a brief."

    me = login(client, "approver@dev.local")
    r = client.post(f"/api/briefs/{bid}/approve")
    assert r.status_code == 200
    body = r.json()
    assert body["approved"] is True
    assert body["approved_by"]["id"] == me["id"]
    assert body["approved_at"]
    assert client.post(f"/api/briefs/{bid}/approve").status_code == 409
    assert client.post("/api/briefs/nope/approve").status_code == 404


def test_approve_refused_on_archived_project(client: TestClient, fake_llm: FakeOpenRouter):
    login(client, ADMIN_EMAIL)
    pid = project_id(client, "Employee Assistant")
    fake_llm.reply(GOOD_BRIEF)
    bid = client.post(f"/api/projects/{pid}/briefs").json()["id"]
    client.patch(f"/api/projects/{pid}", json={"archived": True})
    login(client, "approver@dev.local")
    assert client.post(f"/api/briefs/{bid}/approve").status_code == 409


def test_brief_when_ai_not_set_up(client: TestClient):
    login(client, ADMIN_EMAIL)
    pid = project_id(client, "Employee Assistant")
    r = client.post(f"/api/projects/{pid}/briefs")
    assert r.status_code == 503
    assert r.json()["detail"] == "The AI is not set up yet. Please ask the City AI team."


# ---- Portfolio summary (streamed) ----


def _events(text: str) -> list[tuple[str, str]]:
    out = []
    for block in text.strip().split("\n\n"):
        lines = dict(line.split(": ", 1) for line in block.splitlines())
        out.append((lines["event"], lines["data"]))
    return out


def test_summary_streams_text(client: TestClient, fake_llm: FakeOpenRouter, db: Session):
    login(client, "owner@dev.local")
    fake_llm.stream(["Four projects ", "must fix ", "data first."])
    r = client.post("/api/portfolio/summary")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/event-stream")
    events = _events(r.text)
    assert [e for e, _ in events] == ["text", "text", "text", "done"]
    assert "".join(__import__("json").loads(d)["text"] for e, d in events if e == "text") == (
        "Four projects must fix data first."
    )

    sent = fake_llm.requests[0]["json"]
    assert sent["stream"] is True
    prompt = sent["messages"][0]["content"]
    assert "deadline 30 Oct 2026" in prompt
    assert "Employee Assistant (Hypertrade, Feasibility): Go with actions, weakest" in prompt
    assert "City Care Agent (Retail / CMHL, Production): Not assessed, 0/40 answered." in prompt
    assert prompt.count("/40 answered.") == 12
    assert "Business owner has not set" not in prompt  # no evidence text, only results

    [call] = _calls(db, "summary")
    assert call.ok and call.project_id is None
    assert call.prompt_tokens == 300 and call.cost_usd == 0.0002


def test_summary_errors_are_events_and_logged(
    client: TestClient, fake_llm: FakeOpenRouter, db: Session
):
    login(client, "owner@dev.local")
    fake_llm.stream([], status=500)
    events = _events(client.post("/api/portfolio/summary").text)
    assert events[0][0] == "error"
    assert "did not answer" in events[0][1]

    fake_llm.stream([])  # nothing written
    assert _events(client.post("/api/portfolio/summary").text)[0][0] == "error"

    fake_llm.fail(httpx.ReadTimeout("slow"))
    events = _events(client.post("/api/portfolio/summary").text)
    assert "too long" in events[0][1]
    assert [c.error for c in _calls(db, "summary")] == ["http_500", "empty", "timeout"]


def test_summary_when_ai_not_set_up(client: TestClient):
    login(client, "owner@dev.local")
    r = client.post("/api/portfolio/summary")
    assert r.status_code == 503


# ---- Export ----


def test_export_csv_columns_and_rows(client: TestClient):
    login(client, ADMIN_EMAIL)
    pid = project_id(client, "CFC Model")
    client.put(
        f"/api/projects/{pid}/answers/1.2", json={"answer": "Yes", "evidence": "=HYPERLINK(1)"}
    )
    r = client.get("/api/export.csv")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/csv")
    assert 'filename="city-prism-portfolio-' in r.headers["content-disposition"]
    assert r.content.startswith(b"\xef\xbb\xbf")  # UTF-8 mark for Excel
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert (
        rows[0]
        == COLUMNS
        == [
            "Project",
            "Business unit",
            "Stage",
            "Mode",
            "Dimension",
            "Question ID",
            "Question",
            "Answer",
            "Evidence",
            "Source",
            "Dimension score",
            "Verdict",
        ]
    )
    assert len(rows) == 1 + 12 * 40
    ea = {r[5]: r for r in rows if r[0] == "Employee Assistant"}
    assert ea["5.1"][7:] == [
        "No",
        "Business owner has not set the quality %.",
        "Person",
        "3.0",
        "Go with actions",
    ]
    assert ea["1.1"][7] == "" and ea["1.1"][10] == ""
    cfc = {r[5]: r for r in rows if r[0] == "CFC Model"}
    assert cfc["1.2"][8] == "'=HYPERLINK(1)"  # never run as a formula
    assert ea["5.1"][3] == "Assessing an existing project"


def test_export_xlsx_has_portfolio_sheet(client: TestClient):
    login(client, ADMIN_EMAIL)
    pid = project_id(client, "CFC Model")
    client.put(f"/api/projects/{pid}/answers/1.2", json={"answer": "Yes", "evidence": "=1+1"})
    r = client.get("/api/export.xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Answers", "Portfolio"]
    answers = list(wb["Answers"].iter_rows(values_only=True))
    assert list(answers[0]) == COLUMNS
    assert len(answers) == 1 + 12 * 40
    formula = [row for row in answers if row[0] == "CFC Model" and row[5] == "1.2"][0]
    assert formula[8] == "=1+1"
    assert wb["Answers"].cell(row=answers.index(formula) + 1, column=9).data_type == "s"

    portfolio = list(wb["Portfolio"].iter_rows(values_only=True))
    head = portfolio[0]
    assert head[:4] == ("Project", "Business unit", "Stage", "Mode")
    assert head[-4:] == ("Lowest", "Weakest link", "Verdict", "Answered (of 40)")
    assert len(head) == 4 + 8 + 4
    assert len(portfolio) == 13
    ea = [row for row in portfolio if row[0] == "Employee Assistant"][0]
    assert ea[-2:] == ("Go with actions", 8)
    assert portfolio[-1][-2] == "Not assessed"  # weakest first, not assessed last


def test_export_is_for_admins(client: TestClient):
    for email in ("owner@dev.local", "reviewer@dev.local", "approver@dev.local"):
        login(client, email)
        assert client.get("/api/export.csv").status_code == 403
        assert client.get("/api/export.xlsx").status_code == 403
    client.post("/api/auth/logout")
    assert client.get("/api/export.csv").status_code == 401
