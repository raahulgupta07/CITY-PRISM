"""Portfolio export (SPEC 8): CSV and XLSX, one row per project x question.

Scores and verdicts come from app.scoring, the same rules the screens use.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.models import Answer, Dimension, Project, Question
from app.scoring import TOTAL_QUESTIONS, VERDICT_LABELS, VERDICT_ORDER, ScoreResult, score_project
from app.services.projects import answers_map

ANSWER_LABELS = {"Yes": "Yes", "Partly": "Partly", "No": "No", "Dont_know": "Don't know"}
SOURCE_LABELS = {"person": "Person", "ai_interview": "AI interview", "ai_evidence": "AI evidence"}
MODE_LABELS = {"plan": "Planning a new project", "assess": "Assessing an existing project"}

COLUMNS = [
    "Project",
    "Business unit",
    "Stage",
    "Mode",
    "Dimension",
    "Question ID",
    "Question",
    "Answer",
    "Evidence",
    "Source",
    "Dimension score",
    "Verdict",
]


@dataclass
class ProjectData:
    project: Project
    answers: dict[str, Answer]
    result: ScoreResult


def load(db: Session) -> tuple[list[Dimension], list[Question], list[ProjectData]]:
    """Every project that is not archived, weakest first as on the portfolio."""
    dims = list(db.scalars(select(Dimension).order_by(Dimension.sort_order)).all())
    questions = list(
        db.scalars(
            select(Question)
            .where(Question.active.is_(True))
            .order_by(Question.dimension_id, Question.number)
        ).all()
    )
    projects = db.scalars(select(Project).where(Project.archived.is_(False))).all()
    by_project: dict[str, dict[str, Answer]] = {p.id: {} for p in projects}
    if projects:
        for a in db.scalars(select(Answer).where(Answer.project_id.in_(list(by_project)))).all():
            by_project[a.project_id][a.question_id] = a
    data = [
        ProjectData(p, by_project[p.id], score_project(answers_map(by_project[p.id].values())))
        for p in projects
    ]
    data.sort(
        key=lambda d: (
            VERDICT_ORDER[d.result.verdict],
            d.result.lowest if d.result.lowest is not None else 9,
            d.project.name.lower(),
        )
    )
    return dims, questions, data


def _detail_rows(dims: list[Dimension], questions: list[Question], data: list[ProjectData]):
    titles = {d.id: d.title for d in dims}
    for d in data:
        scores = d.result.dimension_scores
        verdict = VERDICT_LABELS[d.result.verdict]
        for q in questions:
            a = d.answers.get(q.id)
            yield [
                d.project.name,
                d.project.business_unit,
                d.project.stage,
                MODE_LABELS[d.project.mode],
                f"{q.dimension_id}. {titles.get(q.dimension_id, '')}",
                q.id,
                q.text,
                ANSWER_LABELS.get(a.answer, "") if a and a.answer else "",
                a.evidence if a else "",
                SOURCE_LABELS.get(a.source, a.source) if a and a.answer else "",
                scores.get(q.dimension_id),
                verdict,
            ]


def _safe(value: object) -> object:
    """Stop a spreadsheet from running text as a formula (CSV injection)."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def to_csv(db: Session) -> bytes:
    dims, questions, data = load(db)
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(COLUMNS)
    for row in _detail_rows(dims, questions, data):
        writer.writerow(["" if v is None else _safe(v) for v in row])
    # A byte order mark so Excel reads the file as UTF-8.
    return out.getvalue().encode("utf-8-sig")


def _append(ws, row: list[object]) -> None:
    ws.append(row)
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"  # always text, never a formula


def to_xlsx(db: Session) -> bytes:
    dims, questions, data = load(db)
    wb = Workbook()
    detail = wb.active
    detail.title = "Answers"
    _append(detail, COLUMNS)
    for row in _detail_rows(dims, questions, data):
        _append(detail, row)

    portfolio = wb.create_sheet("Portfolio")
    titles = {d.id: d.title for d in dims}
    _append(
        portfolio,
        [
            "Project",
            "Business unit",
            "Stage",
            "Mode",
            *[f"{d.id}. {d.title}" for d in dims],
            "Lowest",
            "Weakest link",
            "Verdict",
            f"Answered (of {TOTAL_QUESTIONS})",
        ],
    )
    for d in data:
        r = d.result
        _append(
            portfolio,
            [
                d.project.name,
                d.project.business_unit,
                d.project.stage,
                MODE_LABELS[d.project.mode],
                *[r.dimension_scores.get(dim.id) for dim in dims],
                r.lowest,
                " and ".join(titles.get(w, str(w)) for w in r.weakest),
                VERDICT_LABELS[r.verdict],
                r.answered,
            ],
        )

    widths = {"Answers": [28, 20, 14, 26, 30, 11, 60, 11, 60, 14, 15, 22]}
    for ws in (detail, portfolio):
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, width in enumerate(widths.get(ws.title, [28, 20, 14, 26] + [12] * 12)):
            ws.column_dimensions[ws.cell(row=1, column=i + 1).column_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
