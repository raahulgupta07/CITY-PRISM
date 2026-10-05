"""Read the text out of an uploaded file, on the server (SPEC 6.3)."""

from __future__ import annotations

import io
from pathlib import PurePath

MAX_FILE_BYTES = 10 * 1024 * 1024
SUPPORTED = (".txt", ".md", ".csv", ".docx", ".pdf")


class ExtractError(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    suffix = PurePath(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise ExtractError("This file type cannot be read. Use a txt, md, csv, docx or pdf file.")
    if len(data) > MAX_FILE_BYTES:
        raise ExtractError("This file is too large. The limit is 10 MB.")
    try:
        if suffix in (".txt", ".md", ".csv"):
            text = data.decode("utf-8-sig", errors="replace")
        elif suffix == ".docx":
            text = _docx(data)
        else:
            text = _pdf(data)
    except ExtractError:
        raise
    except Exception as exc:  # broken or protected files
        raise ExtractError("This file could not be read. Try pasting the text instead.") from exc
    text = "\n".join(line.rstrip() for line in text.replace("\r\n", "\n").split("\n")).strip()
    if not text:
        raise ExtractError("No text was found in this file. Try pasting the text instead.")
    return text


def _docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ExtractError("This PDF is protected. Try pasting the text instead.")
    return "\n".join(page.extract_text() or "" for page in reader.pages)
